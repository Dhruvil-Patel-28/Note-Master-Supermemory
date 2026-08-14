import io

import openpyxl
from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv", ".json", ".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp"}

TEXT_LAYER_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".docx", ".xlsx", ".pdf"}


def has_text_layer(path, content_type: str = "") -> bool:
    ext = path.suffix.lower()
    if ext not in TEXT_LAYER_EXTENSIONS:
        return False
    if ext == ".pdf":
        try:
            return any((page.extract_text() or "").strip() for page in PdfReader(path).pages)
        except Exception:
            return False
    return True


def extract_text(path, ext: str = "") -> str:
    ext = ext or path.suffix.lower()
    if ext in {".txt", ".md", ".csv", ".json"}:
        return path.read_text(errors="replace")
    if ext == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if ext == ".docx":
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join("" if c is None else str(c) for c in row))
        return "\n".join(rows)
    raise ValueError(f"unsupported extension: {ext}")


def sniff_text(data: bytes) -> str:
    return io.TextIOWrapper(io.BytesIO(data), encoding="utf-8", errors="replace").read()